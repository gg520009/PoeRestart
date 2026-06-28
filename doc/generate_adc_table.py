import math
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

# Parameters for SDNT1608X103F3380FTF
R25 = 10000.0  # 10k ohms
B = 3380.0     # B-constant
T25 = 298.15   # 25°C in Kelvin

# Circuit parameters
VCC = 3.3       # 3.3V
R_PU = 82000.0  # 82k ohms pull-up
ADC_MAX = 4095  # 12-bit ADC (0 to 4095)

# Calculate resistance from temperature
def calc_resistance(temp_c):
    tk = temp_c + 273.15
    r = R25 * math.exp(B * (1.0 / tk - 1.0 / T25))
    return r

# Calculate temperature from resistance
def calc_temperature(r_ntc):
    tk = 1.0 / (1.0 / T25 + (1.0 / B) * math.log(r_ntc / R25))
    return tk - 273.15

# Create workbook
wb = openpyxl.Workbook()

# ==========================================
# Sheet 1: NTC ADC Table (Temp -> ADC)
# ==========================================
ws1 = wb.active
ws1.title = "NTC ADC Table"
ws1.views.sheetView[0].showGridLines = True

# Title block
ws1.merge_cells("A1:F1")
ws1["A1"] = "SDNT1608X103F3380FTF 热敏电阻 12位ADC采样对照表"
ws1["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="FFFFFF")
ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws1["A1"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ws1.row_dimensions[1].height = 45

# Spec info block
specs = [
    ("产品型号 (Part Number)", "SDNT1608X103F3380FTF", "供电电压 (VCC)", "3.3 V"),
    ("标称阻值 (R25)", "10 kΩ ± 1%", "上拉电阻 (R_pullup)", "82 kΩ"),
    ("B常数 (B25/50)", "3380 K ± 1%", "ADC 分辨率", "12-bit (0 ~ 4095)"),
    ("工作温度范围", "-55°C ~ +125°C", "采样电路结构", "VCC -> R_pu -> ADC Pin / NTC -> GND"),
]

thin_side = Side(style='thin', color='D9D9D9')
font_bold = Font(name="Microsoft YaHei", size=9, bold=True)
font_regular = Font(name="Microsoft YaHei", size=9)
fill_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

for idx, (param1, val1, param2, val2) in enumerate(specs, start=2):
    ws1.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=2)
    cell_p1 = ws1.cell(row=idx, column=1, value=param1)
    cell_v1 = ws1.cell(row=idx, column=3, value=val1)
    
    ws1.merge_cells(start_row=idx, start_column=4, end_row=idx, end_column=5)
    cell_p2 = ws1.cell(row=idx, column=4, value=param2)
    cell_v2 = ws1.cell(row=idx, column=6, value=val2)
    
    cell_p1.font = font_bold
    cell_p1.fill = fill_gray
    cell_v1.font = font_regular
    cell_p2.font = font_bold
    cell_p2.fill = fill_gray
    cell_v2.font = font_regular
    
    cell_p1.alignment = Alignment(horizontal="left", vertical="center")
    cell_v1.alignment = Alignment(horizontal="left", vertical="center")
    cell_p2.alignment = Alignment(horizontal="left", vertical="center")
    cell_v2.alignment = Alignment(horizontal="left", vertical="center")
    
    for c in range(1, 7):
        ws1.cell(row=idx, column=c).border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    ws1.row_dimensions[idx].height = 20

ws1.row_dimensions[6].height = 15

# Table headers
headers1 = ["温度 (°C)", "NTC电阻 (Ω)", "NTC电阻 (kΩ)", "采样电压 (V)", "ADC十进制值", "ADC十六进制值"]
for col_idx, header in enumerate(headers1, start=1):
    cell = ws1.cell(row=7, column=col_idx, value=header)
    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        top=Side(style='medium', color='1F4E79'),
        bottom=Side(style='medium', color='1F4E79'),
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9')
    )
ws1.row_dimensions[7].height = 25

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

row_idx = 8
for temp_c in range(-55, 126):
    r = calc_resistance(temp_c)
    r_k = r / 1000.0
    v_adc = VCC * r / (R_PU + r)
    adc_dec = round(r / (R_PU + r) * ADC_MAX)
    adc_hex = f"0x{adc_dec:03X}"
    
    c_temp = ws1.cell(row=row_idx, column=1, value=temp_c)
    c_r = ws1.cell(row=row_idx, column=2, value=r)
    c_rk = ws1.cell(row=row_idx, column=3, value=r_k)
    c_v = ws1.cell(row=row_idx, column=4, value=v_adc)
    c_dec = ws1.cell(row=row_idx, column=5, value=adc_dec)
    c_hex = ws1.cell(row=row_idx, column=6, value=adc_hex)
    
    c_temp.number_format = '0'
    c_r.number_format = '#,##0.0'
    c_rk.number_format = '0.000'
    c_v.number_format = '0.0000'
    c_dec.number_format = '0'
    
    c_temp.alignment = Alignment(horizontal="center", vertical="center")
    c_r.alignment = Alignment(horizontal="right", vertical="center")
    c_rk.alignment = Alignment(horizontal="right", vertical="center")
    c_v.alignment = Alignment(horizontal="right", vertical="center")
    c_dec.alignment = Alignment(horizontal="right", vertical="center")
    c_hex.alignment = Alignment(horizontal="center", vertical="center")
    
    for col in range(1, 7):
        cell = ws1.cell(row=row_idx, column=col)
        cell.font = Font(name="Microsoft YaHei", size=9)
        cell.border = thin_border
        if temp_c % 2 == 0:
            cell.fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
            
    ws1.row_dimensions[row_idx].height = 20
    row_idx += 1

for col in ws1.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col[6:]:
        if cell.value is not None:
            if isinstance(cell.value, float):
                val_str = f"{cell.value:,.4f}"
            else:
                val_str = str(cell.value)
            max_len = max(max_len, len(val_str))
    ws1.column_dimensions[col_letter].width = max(max_len + 8, 15)


# ==========================================
# Sheet 2: ADC to Temp Table (ADC step -1 -> Temp)
# ==========================================
ws2 = wb.create_sheet(title="ADC to Temp Table")
ws2.views.sheetView[0].showGridLines = True

# Title block
ws2.merge_cells("A1:F1")
ws2["A1"] = "SDNT1608X103F3380FTF 12位ADC采样值转温度对照表 (-50°C ~ 125°C, ADC递减步长 1)"
ws2["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="FFFFFF")
ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws2["A1"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ws2.row_dimensions[1].height = 45

# Spec info block
for idx, (param1, val1, param2, val2) in enumerate(specs, start=2):
    ws2.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=2)
    cell_p1 = ws2.cell(row=idx, column=1, value=param1)
    cell_v1 = ws2.cell(row=idx, column=3, value=val1)
    
    ws2.merge_cells(start_row=idx, start_column=4, end_row=idx, end_column=5)
    cell_p2 = ws2.cell(row=idx, column=4, value=param2)
    cell_v2 = ws2.cell(row=idx, column=6, value=val2)
    
    cell_p1.font = font_bold
    cell_p1.fill = fill_gray
    cell_v1.font = font_regular
    cell_p2.font = font_bold
    cell_p2.fill = fill_gray
    cell_v2.font = font_regular
    
    cell_p1.alignment = Alignment(horizontal="left", vertical="center")
    cell_v1.alignment = Alignment(horizontal="left", vertical="center")
    cell_p2.alignment = Alignment(horizontal="left", vertical="center")
    cell_v2.alignment = Alignment(horizontal="left", vertical="center")
    
    for c in range(1, 7):
        ws2.cell(row=idx, column=c).border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    ws2.row_dimensions[idx].height = 20

ws2.row_dimensions[6].height = 15

# Table headers for Sheet 2
headers2 = ["ADC十进制值", "ADC十六进制值", "采样电压 (V)", "NTC电阻 (Ω)", "NTC电阻 (kΩ)", "对应摄氏温度 (°C)"]
for col_idx, header in enumerate(headers2, start=1):
    cell = ws2.cell(row=7, column=col_idx, value=header)
    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        top=Side(style='medium', color='1F4E79'),
        bottom=Side(style='medium', color='1F4E79'),
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9')
    )
ws2.row_dimensions[7].height = 25

# ADC values from 3466 down to 29 (covering -50°C to 125°C)
row_idx = 8
for adc_dec in range(3466, 28, -1):
    v_adc = VCC * adc_dec / ADC_MAX
    r_ntc = (adc_dec * R_PU) / (ADC_MAX - adc_dec)
    r_k = r_ntc / 1000.0
    temp_c = calc_temperature(r_ntc)
    adc_hex = f"0x{adc_dec:03X}"
    
    c_dec = ws2.cell(row=row_idx, column=1, value=adc_dec)
    c_hex = ws2.cell(row=row_idx, column=2, value=adc_hex)
    c_v = ws2.cell(row=row_idx, column=3, value=v_adc)
    c_r = ws2.cell(row=row_idx, column=4, value=r_ntc)
    c_rk = ws2.cell(row=row_idx, column=5, value=r_k)
    c_temp = ws2.cell(row=row_idx, column=6, value=temp_c)
    
    c_dec.number_format = '0'
    c_v.number_format = '0.0000'
    c_r.number_format = '#,##0.0'
    c_rk.number_format = '0.000'
    c_temp.number_format = '0.00'
    
    c_dec.alignment = Alignment(horizontal="right", vertical="center")
    c_hex.alignment = Alignment(horizontal="center", vertical="center")
    c_v.alignment = Alignment(horizontal="right", vertical="center")
    c_r.alignment = Alignment(horizontal="right", vertical="center")
    c_rk.alignment = Alignment(horizontal="right", vertical="center")
    c_temp.alignment = Alignment(horizontal="right", vertical="center")
    
    for col in range(1, 7):
        cell = ws2.cell(row=row_idx, column=col)
        cell.font = Font(name="Microsoft YaHei", size=9)
        cell.border = thin_border
        if (3466 - adc_dec) % 2 == 0:
            cell.fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
            
    ws2.row_dimensions[row_idx].height = 20
    row_idx += 1

for col in ws2.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col[6:]:
        if cell.value is not None:
            if isinstance(cell.value, float):
                val_str = f"{cell.value:,.4f}"
            else:
                val_str = str(cell.value)
            max_len = max(max_len, len(val_str))
    ws2.column_dimensions[col_letter].width = max(max_len + 8, 15)

# Save the workbook to doc folder
script_dir = os.path.dirname(os.path.abspath(__file__))
output_file = os.path.join(script_dir, "SDNT1608X103F3380FTF_ADC_Table.xlsx")
wb.save(output_file)
print(f"Excel sheet '{output_file}' generated successfully!")
