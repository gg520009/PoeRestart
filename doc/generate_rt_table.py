import math
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Parameters for SDNT1608X103F3380FTF
R25 = 10000.0  # 10k ohms
B = 3380.0     # B-constant
T25 = 298.15   # 25°C in Kelvin

# Calculate resistance
def calc_resistance(temp_c):
    tk = temp_c + 273.15
    r = R25 * math.exp(B * (1.0 / tk - 1.0 / T25))
    return r

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "NTC R-T Table"

# Grid lines visible
ws.views.sheetView[0].showGridLines = True

# Title block
ws.merge_cells("A1:D1")
ws["A1"] = "SDNT1608X103F3380FTF 阻值-温度对照表 (R-T Table)"
ws["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="FFFFFF")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws["A1"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ws.row_dimensions[1].height = 40

# Spec info block
specs = [
    ("产品型号 (Part Number)", "SDNT1608X103F3380FTF"),
    ("标称阻值 (R25)", "10 kΩ ± 1%"),
    ("B常数 (B25/50)", "3380 K ± 1%"),
    ("工作温度范围 (Operating Temp)", "-55°C ~ +125°C"),
]
for idx, (param, val) in enumerate(specs, start=2):
    ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=2)
    ws.merge_cells(start_row=idx, start_column=3, end_row=idx, end_column=4)
    
    # Set values
    cell_p = ws.cell(row=idx, column=1, value=param)
    cell_v = ws.cell(row=idx, column=3, value=val)
    
    # Styling
    cell_p.font = Font(name="Microsoft YaHei", size=10, bold=True)
    cell_p.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    cell_v.font = Font(name="Microsoft YaHei", size=10)
    
    # Alignments
    cell_p.alignment = Alignment(horizontal="left", vertical="center")
    cell_v.alignment = Alignment(horizontal="left", vertical="center")
    
    # Borders for spec block
    thin_side = Side(style='thin', color='D9D9D9')
    for c in range(1, 5):
        ws.cell(row=idx, column=c).border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        
    ws.row_dimensions[idx].height = 20

# Blank row
ws.row_dimensions[6].height = 15

# Table headers
headers = ["温度 (°C)", "绝对温度 (K)", "电阻值 (Ω)", "电阻值 (kΩ)"]
for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=7, column=col_idx, value=header)
    cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=7, column=col_idx).border = Border(
        top=Side(style='medium', color='1F4E79'),
        bottom=Side(style='medium', color='1F4E79'),
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9')
    )
ws.row_dimensions[7].height = 25

# Border styles for data cells
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# Fill data
row_idx = 8
for temp_c in range(-55, 126):
    r = calc_resistance(temp_c)
    tk = temp_c + 273.15
    r_k = r / 1000.0
    
    # Cells
    c_temp = ws.cell(row=row_idx, column=1, value=temp_c)
    c_tk = ws.cell(row=row_idx, column=2, value=tk)
    c_r = ws.cell(row=row_idx, column=3, value=r)
    c_rk = ws.cell(row=row_idx, column=4, value=r_k)
    
    # Number formats
    c_temp.number_format = '0'
    c_tk.number_format = '0.00'
    c_r.number_format = '#,##0.0'
    c_rk.number_format = '0.000'
    
    # Alignments
    c_temp.alignment = Alignment(horizontal="center", vertical="center")
    c_tk.alignment = Alignment(horizontal="center", vertical="center")
    c_r.alignment = Alignment(horizontal="right", vertical="center")
    c_rk.alignment = Alignment(horizontal="right", vertical="center")
    
    # Styling and borders
    for col in range(1, 5):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = Font(name="Microsoft YaHei", size=10)
        cell.border = thin_border
        
        # Alternate row coloring (light blue/gray)
        if temp_c % 2 == 0:
            cell.fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
            
    ws.row_dimensions[row_idx].height = 20
    row_idx += 1

# Autofit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    # Check lengths from row 7 downwards (table data)
    for cell in col[6:]:
        if cell.value is not None:
            # Format display length approximation
            if isinstance(cell.value, float):
                val_str = f"{cell.value:,.2f}"
            else:
                val_str = str(cell.value)
            max_len = max(max_len, len(val_str))
    ws.column_dimensions[col_letter].width = max(max_len + 8, 16)

# Save the workbook to doc folder
import os
script_dir = os.path.dirname(os.path.abspath(__file__))
output_file = os.path.join(script_dir, "SDNT1608X103F3380FTF_RT_Table.xlsx")
wb.save(output_file)
print(f"Excel sheet '{output_file}' generated successfully!")
